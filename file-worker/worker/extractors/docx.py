import asyncio
import io
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from worker.extractors.base import (
    ExtractionError,
    escape_table_cell,
    rows_to_markdown_table,
)

_HEADING_STYLE = re.compile(r"^Heading (\d+)$")


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style else None
    if not style_name:
        return None
    if style_name == "Title":
        return 1
    match = _HEADING_STYLE.match(style_name)
    return int(match.group(1)) if match else None


def _paragraph_to_markdown(paragraph: Paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    level = _heading_level(paragraph)
    if level:
        return f"{'#' * min(level, 6)} {text}"
    return text


def _table_to_markdown(table: Table) -> str:
    rows = [[escape_table_cell(cell.text) for cell in row.cells] for row in table.rows]
    return rows_to_markdown_table(rows)


def _extract_sync(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"could not parse DOCX: {exc}") from exc

    # document.paragraphs/.tables are separate flat lists that don't
    # preserve document order (a table would always end up after every
    # paragraph regardless of where it actually sits) - walking the raw
    # body elements keeps paragraphs and tables interleaved correctly.
    blocks = []
    for child in document.element.body:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _paragraph_to_markdown(Paragraph(child, document))
        elif tag == "tbl":
            text = _table_to_markdown(Table(child, document))
        else:
            continue
        if text:
            blocks.append(text)

    return "\n\n".join(blocks)


async def extract(name: str, data: bytes) -> str:
    content = await asyncio.to_thread(_extract_sync, data)
    return content

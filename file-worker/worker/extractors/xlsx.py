import asyncio
import io

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from worker.extractors.base import (
    ExtractionError,
    escape_table_cell,
    rows_to_markdown_table,
)


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return escape_table_cell(str(value))


def _sheet_to_markdown(sheet: Worksheet) -> str:
    rows = [
        [_cell_to_str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    table = rows_to_markdown_table(rows)
    if not table:
        return ""
    return f"# {sheet.title}\n\n{table}"


def _extract_sync(data: bytes) -> str:
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data),
            data_only=True,
            read_only=True,
        )
    except Exception as exc:
        raise ExtractionError(f"could not parse XLSX: {exc}") from exc

    sheets = [_sheet_to_markdown(workbook[name]) for name in workbook.sheetnames]
    return "\n\n".join(sheet for sheet in sheets if sheet)


async def extract(name: str, data: bytes) -> str:
    return await asyncio.to_thread(_extract_sync, data)

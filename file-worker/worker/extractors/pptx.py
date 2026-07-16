import asyncio
import io

from pptx import Presentation
from pptx.slide import Slide
from pptx.table import Table

from worker.extractors.base import (
    ExtractionError,
    escape_table_cell,
    rows_to_markdown_table,
)


def _table_to_markdown(table: Table) -> str:
    rows = [[escape_table_cell(cell.text) for cell in row.cells] for row in table.rows]
    return rows_to_markdown_table(rows)


def _slide_to_markdown(slide: Slide) -> str:
    blocks = []

    title = slide.shapes.title
    if title is not None and title.has_text_frame:
        text = title.text_frame.text.strip()
        if text:
            blocks.append(f"# {text}")

    for shape in slide.shapes:
        if shape == title:
            continue
        if shape.has_table:
            text = _table_to_markdown(shape.table)  # type: ignore[attr-defined]
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()  # type: ignore[attr-defined]
        else:
            continue
        if text:
            blocks.append(text)

    return "\n\n".join(blocks)


def _extract_sync(data: bytes) -> str:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"could not parse PPTX: {exc}") from exc

    slides = [_slide_to_markdown(slide) for slide in presentation.slides]
    return "\n\n".join(block for block in slides if block)


async def extract(name: str, data: bytes) -> str:
    return await asyncio.to_thread(_extract_sync, data)
